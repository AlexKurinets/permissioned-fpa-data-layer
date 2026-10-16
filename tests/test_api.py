import io
import time

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from pptx import Presentation

from conftest import ANALYST, CFO, FPA_LEAD, GM_NEWS, PLATFORM_MGR, VP_ENG
from fpa.config import get_settings


def _deck_text(content: bytes) -> str:
    prs = Presentation(io.BytesIO(content))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts += [cell.text for row in shape.table.rows for cell in row.cells]
    return "\n".join(parts)


def test_health_is_open(api):
    body = api.get("/health").json()
    assert body["status"] == "ok" and body["latest_run"]["status"] == "published"


def _forged(**claims):
    s = get_settings()
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    now = int(time.time())
    base = {"iss": s.oidc_issuer, "aud": s.oidc_audience, "email": CFO, "iat": now, "exp": now + 600}
    base.update(claims)
    return jwt.encode(base, key, algorithm="RS256")


def test_tokens_are_checked(api):
    assert api.get("/me").status_code == 401
    assert api.get("/me", headers={"Authorization": f"Bearer {_forged()}"}).status_code == 401
    issuer = api.app.state.issuer
    expired = issuer.issue(CFO, ttl_seconds=-120)
    wrong_audience = issuer.issue(CFO, aud="some-other-app")
    for token in (expired, wrong_audience):
        assert api.get("/me", headers={"Authorization": f"Bearer {token}"}).status_code == 401
    unknown = issuer.issue("stranger@example.com")
    assert api.get("/me", headers={"Authorization": f"Bearer {unknown}"}).status_code == 403


def test_me_reports_scope(api, token):
    me = api.get("/me", headers=token(VP_ENG)).json()
    assert me["cost_centers"] == ["4110", "4120", "4130"]
    assert "department:CORP-ENG" in me["units"] and me["permissions"] == []


def test_bva_is_scoped(api, token):
    eng = api.get("/bva?period=2026-09", headers=token(VP_ENG)).json()
    summary = {m["key"]: m for m in eng["summary"]}
    assert summary["net_revenue"]["in_scope"] is False
    assert summary["personnel_cost"]["in_scope"] is True
    assert summary["adjusted_ebitda"]["in_scope"] is False  # no revenue in scope, so no margin
    assert {line["unit_name"] for line in eng["lines"]} == {"Engineering"}
    platform = api.get("/bva?period=2026-09", headers=token(PLATFORM_MGR)).json()
    assert {(line["unit_type"], line["unit_name"]) for line in platform["lines"]} == {
        ("cost_center", "Platform engineering")}


def test_variances_and_review_permissions(api, token, owner):
    assert len(api.get("/variances", headers=token(CFO)).json()["lines"]) == 10
    eng = api.get("/variances", headers=token(VP_ENG)).json()["lines"]
    assert len(eng) == 2
    target = eng[0]["variance_id"]
    assert api.post(f"/variances/{target}/review", json={"action": "accept"}, headers=token(VP_ENG)).status_code == 403
    assert api.post(f"/variances/{target}/review", json={"action": "edit"}, headers=token(ANALYST)).status_code == 422
    ok = api.post(f"/variances/{target}/review", json={"action": "accept"}, headers=token(ANALYST))
    assert ok.status_code == 200 and ok.json()["status"] == "accepted"
    owner.execute("update core.variance_line set status = 'draft', final_text = null, reviewed_by = null,"
                  " reviewed_at = null, review_note = null where variance_id = %s", (target,))


@pytest.mark.parametrize("question, person, expected", [
    ("What did we spend on contractors in Engineering in September?", VP_ENG,
     "Contractors (6050) for Engineering, September 2026: $80,100 actual against $36,000 budget, $44,100 over"),
    ("display advertising in Local News vs budget", VP_ENG, "You don't have access to Local News."),
    ("salaries in Engineering", PLATFORM_MGR, "You can see 1 of the 3 cost centers in Engineering"),
    ("net revenue year to date", CFO, "Net revenue for everything you have access to, year to date through"),
    ("digital subscriptions in Q3 vs last year", GM_NEWS,
     "Digital subscriptions (4110) for everything you have access to, Q3 2026"),
])
def test_ask_answers_inside_the_askers_scope(api, token, question, person, expected):
    body = api.post("/ask", json={"question": question}, headers=token(person)).json()
    assert body["answer"].startswith(expected), body["answer"]


def test_ask_explains_when_it_cannot_map_a_question(api, token):
    body = api.post("/ask", json={"question": "how is morale?"}, headers=token(CFO)).json()
    assert "semantic layer" in body["answer"] and body["query"] is None


def test_package_contains_only_the_callers_scope(api, token):
    eng = api.get("/packages/2026-09.pptx", headers=token(VP_ENG))
    assert eng.status_code == 200
    text = _deck_text(eng.content)
    assert "Engineering" in text and "Contractors" in text
    assert "Local Ad Sales" not in text and "Event Operations" not in text
    cfo = _deck_text(api.get("/packages/2026-09.pptx", headers=token(CFO)).content)
    assert "Local Ad Sales" in cfo and "Event Operations" in cfo and "Scope: Company" in cfo


def test_headcount_forecast_and_definitions_endpoints(api, token):
    hc = api.get("/headcount", headers=token(FPA_LEAD)).json()
    assert hc["salary_visible"] is True and len(hc["exceptions"]) == 3
    fc = api.get("/forecast", headers=token(GM_NEWS)).json()
    assert [r["unit_id"] for r in fc["revenue"]] == ["LN"]
    accounts = api.get("/semantic/accounts", headers=token(PLATFORM_MGR)).json()
    assert len(accounts) == 63


def test_access_review_endpoint(api, token):
    assert api.get("/admin/access-review", headers=token(VP_ENG)).status_code == 403
    assert len(api.get("/admin/access-review", headers=token(FPA_LEAD)).json()) >= 12
