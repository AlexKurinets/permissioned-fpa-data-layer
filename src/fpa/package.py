"""The monthly reporting package, generated per person from the same scoped queries.

A department head and the CFO request the same package and get different decks,
because every number on every slide comes from a query that ran under their grants.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime

import psycopg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from fpa import headcount, reporting
from fpa.semantic import Catalog
from fpa.variance import money, month_label

NAVY = RGBColor(0x1F, 0x2A, 0x44)
INK = RGBColor(0x26, 0x2B, 0x33)
MUTED = RGBColor(0x6B, 0x72, 0x80)
GOOD = RGBColor(0x1B, 0x7F, 0x5A)
BAD = RGBColor(0xB4, 0x3C, 0x2F)
RULE = RGBColor(0xD9, 0xDE, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class PackageBlocked(RuntimeError):
    pass


def scope_label(me: dict) -> str:
    grants = me["grants"]
    if any(g["scope_type"] == "all" for g in grants):
        return "Company"
    return ", ".join(f"{g['scope_type'].replace('_', ' ')} {g['scope_value']}" for g in grants) or "No access"


def _text(slide, left, top, width, height, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def _header(prs, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _text(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6), title, 26, True, NAVY)
    _text(slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.4), subtitle, 13, False, MUTED)
    return slide


def _table(slide, rows: list[list[str]], left, top, width, col_widths: list[float], colors=None, size=11):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, Inches(0.3) * len(rows))
    table = shape.table
    total = sum(col_widths)
    for i, w in enumerate(col_widths):
        table.columns[i].width = Emu(int(width * w / total))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else RGBColor(0xF4, 0xF6, 0xF9))
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(size)
                para.font.name = "Calibri"
                para.font.bold = r == 0
                para.alignment = PP_ALIGN.RIGHT if c > 0 and r > 0 and _numeric(value) else PP_ALIGN.LEFT
                para.font.color.rgb = WHITE if r == 0 else (colors(r, c) if colors else INK) or INK
    return table


def _numeric(value: str) -> bool:
    return bool(value) and (value[0] in "$+-(" or value[0].isdigit() or value == "n/a")


def _signed(value: float) -> str:
    return ("+" if value >= 0 else "-") + money(value)


def build(conn: psycopg.Connection, catalog: Catalog, period: date, version: str, me: dict) -> bytes:
    status = reporting.status(conn)
    published = [p for p in reporting.periods(conn)]
    latest = status["latest_run"]
    if latest and latest["status"] != "published":
        raise PackageBlocked(f"The latest refresh ({latest['run_id']}) ended with {latest['status']}. "
                             "The package is only built from data that passed the tie-out.")
    if period not in published:
        raise PackageBlocked(f"{month_label(period)} has not been published.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    scope = scope_label(me)
    variances = reporting.variance_list(conn, period)
    reviewed = sum(1 for v in variances if v["status"] != "draft")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(3.2))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    _text(slide, Inches(0.8), Inches(1.0), Inches(11.5), Inches(1.0),
          f"{month_label(period)} budget vs actual", 40, True, WHITE)
    _text(slide, Inches(0.8), Inches(2.0), Inches(11.5), Inches(0.6),
          f"First-pass package for {me['full_name']}, {me['title']}  |  Scope: {scope}", 18, False, WHITE)
    run = status["latest_published_run"]
    lines = [
        f"Data: refresh {run['run_id']}, the GL tied to the ERP trial balance for every entity and period.",
        f"Commentary: {reviewed} of {len(variances)} material lines reviewed by FP&A. "
        "Unreviewed lines are marked Draft.",
        f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC from the permissioned data layer. "
        "Every figure reflects only the data this person is granted.",
    ]
    for i, line in enumerate(lines):
        _text(slide, Inches(0.8), Inches(3.8 + 0.6 * i), Inches(11.8), Inches(0.5), line, 15, False, INK)

    summary = [m for m in reporting.bva_summary(conn, catalog, period, version) if m["in_scope"]]
    slide = _header(prs, "Summary", f"{month_label(period)} and year to date, {scope.lower()} scope")
    rows = [["Metric", "Actual", "Budget", "Variance", "Var %", "YTD actual", "YTD budget", "YTD variance"]]
    for m in summary:
        mo, ytd = m["month"], m["ytd"]
        rows.append([m["label"], money(mo["actual"]), money(mo["budget"]), _signed(mo["variance"]),
                     f"{mo['variance_pct']:+.1f}%" if mo["variance_pct"] is not None else "n/a",
                     money(ytd["actual"]), money(ytd["budget"]), _signed(ytd["variance"])])
    favorable = {i + 1: (m["month"]["favorable"], m["ytd"]["favorable"]) for i, m in enumerate(summary)}
    if len(rows) == 1:
        _text(slide, Inches(0.6), Inches(1.6), Inches(12), Inches(0.5), "No budgeted lines in this scope.", 16)
    else:
        _table(slide, rows, Inches(0.6), Inches(1.6), Inches(12.1), [2.6, 1.4, 1.4, 1.4, 1.0, 1.5, 1.5, 1.5],
               colors=lambda r, c: ((GOOD if favorable[r][0] else BAD) if c in (3, 4)
                                    else (GOOD if favorable[r][1] else BAD) if c == 7 else INK))

    chunks = [variances[i:i + 4] for i in range(0, len(variances), 4)] or [[]]
    for n, chunk in enumerate(chunks, start=1):
        suffix = f" ({n} of {len(chunks)})" if len(chunks) > 1 else ""
        slide = _header(prs, f"Material variances{suffix}",
                        f"Lines over {money(catalog.min_abs_variance)} and {catalog.min_pct_variance:.0%} "
                        "of budget, largest first")
        if not chunk:
            _text(slide, Inches(0.6), Inches(1.6), Inches(12), Inches(0.5),
                  "No material variances in the departments this scope covers in full.", 16)
            continue
        rows = [["Department and line", "Variance", "Commentary"]]
        for v in chunk:
            text = v["final_text"] if v["status"] in ("accepted", "edited") else f"Draft: {v['draft']}"
            rows.append([f"{v['department_name']}\n{v['account_name']}", _signed(v["variance"]), text])
        _table(slide, rows, Inches(0.6), Inches(1.6), Inches(12.1), [2.3, 1.2, 8.6],
               colors=lambda r, c, chunk=chunk: (GOOD if chunk[r - 1]["favorable"] else BAD) if c == 1 else INK,
               size=11)

    hc = headcount.report(conn, period, version)
    slide = _header(prs, "Headcount", f"Plan vs actual at the end of {month_label(period)}")
    depts: dict[str, dict[str, int]] = {}
    for r in hc["by_cost_center"]:
        d = depts.setdefault(r["department_name"], {"plan": 0.0, "actual": 0, "open": 0})
        d["plan"] += r["plan"]
        d["actual"] += r["actual"]
        d["open"] += r["open_roles"]
    rows = [["Department", "Plan", "Actual", "Gap", "Open roles"]]
    rows += [[name, f"{d['plan']:.0f}", str(d["actual"]), f"{round(d['actual'] - d['plan']):+d}", str(d["open"])]
             for name, d in sorted(depts.items())]
    _table(slide, rows, Inches(0.6), Inches(1.5), Inches(7.6), [3.2, 1, 1, 1, 1.4], size=10)
    _text(slide, Inches(8.6), Inches(1.6), Inches(4.2), Inches(2.5),
          f"{len(hc['exceptions'])} payroll-vs-HRIS exception(s) to clear before payroll closes. "
          "Details are on the Headcount page of the dashboard.", 14, False, INK)

    fc = reporting.forecast_view(conn, catalog, period, version)
    if fc["revenue"] or fc["opex"]:
        slide = _header(prs, "Full-year outlook",
                        f"Year-to-date actuals through {month_label(period)} plus the rolling forecast")
        rows = [["Unit", "Outlook", "Budget", "Outlook vs budget"]]
        rows += [[f"{r['unit_name']} revenue", money(r["outlook"]), money(r["fy_budget"]),
                  _signed(r["outlook_vs_budget"])] for r in fc["revenue"]]
        rows += [[f"{r['department_name']} opex", money(r["outlook"]), money(r["fy_budget"]),
                  _signed(r["outlook_vs_budget"])] for r in fc["opex"]]
        _table(slide, rows, Inches(0.6), Inches(1.5), Inches(9.0), [3.6, 1.8, 1.8, 1.8], size=10)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
